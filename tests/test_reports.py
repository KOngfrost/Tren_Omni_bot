"""Тесты отчётов: генерация Excel, маскировка анонимных, защита от дублей."""

from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from core.models import Department, Ticket, TicketStatus, User
from core.reporting import DEFAULT_DEPTS, build_daily_report


def _make_data(tickets):
    return {
        "tickets": tickets,
        "departments": [Department(id=1, name="Жилбыт"), Department(id=2, name="Информ")],
    }


def test_build_daily_report_sheets_match_departments():
    """Ошибка #16: количество листов отделов = количеству отделов в БД.

    В данных 2 отдела -> 1 лист сводки + 2 листа отделов (без добивки
    DEFAULT_DEPTS до четырёх).
    """
    data = _make_data([])
    report = build_daily_report(data, datetime(2026, 8, 30))

    workbook = load_workbook(BytesIO(report))
    assert len(workbook.sheetnames) == 3
    assert workbook.sheetnames[0] == "Краткая информация"
    assert "Жилбыт" in workbook.sheetnames
    assert "Информ" in workbook.sheetnames


def test_build_daily_report_dynamic_departments():
    """Ошибка #16: отделов может быть больше четырёх — все попадают в отчёт.

    Критерий приёмки: отделы, добавленные администратором через панель
    управления, получают собственные вкладки в ежедневном Excel-отчёте.
    """
    real_depts = [Department(id=i, name=f"Отдел №{i}") for i in range(1, 7)]
    data = {"tickets": [], "departments": real_depts}
    report = build_daily_report(data, datetime(2026, 8, 30))

    workbook = load_workbook(BytesIO(report))
    assert len(workbook.sheetnames) == 7  # сводка + 6 отделов
    assert workbook.sheetnames[0] == "Краткая информация"
    for i in range(1, 7):
        assert f"Отдел №{i}" in workbook.sheetnames


def test_build_daily_report_falls_back_to_default_depts_when_db_empty():
    """Пустая БД (нет отделов) — dev-фоллбэк DEFAULT_DEPTS (5 листов)."""
    data = {"tickets": [], "departments": []}
    report = build_daily_report(data, datetime(2026, 8, 30))

    workbook = load_workbook(BytesIO(report))
    assert len(workbook.sheetnames) == 5  # сводка + 4 дефолтных отдела
    assert set(workbook.sheetnames[1:]) == set(DEFAULT_DEPTS)


def test_build_daily_report_with_real_db_departments():
    real_depts = [
        Department(id=1, name="Культурно-массовый"),
        Department(id=2, name="Корпоративный"),
        Department(id=3, name="Жилищно-бытовой"),
        Department(id=4, name="Информационный"),
    ]
    data = {"tickets": [], "departments": real_depts}
    report = build_daily_report(data, datetime(2026, 8, 30))

    workbook = load_workbook(BytesIO(report))
    assert len(workbook.sheetnames) == 5
    assert workbook.sheetnames[0] == "Краткая информация"
    assert set(workbook.sheetnames[1:]) == {
        "Культурно-массовый",
        "Корпоративный",
        "Жилищно-бытовой",
        "Информационный",
    }


def test_report_masks_anonymous_user_data():
    ticket = Ticket(
        id=1,
        topic="Анонимная тема",
        description="Текст",
        status=TicketStatus.ANONYMOUS,
        is_anonymous=True,
    )
    ticket.user = User(full_name="Секретное Имя", dormitory="Общежитие №1")
    ticket.department = Department(id=1, name="Жилбыт")

    data = _make_data([ticket])
    report = build_daily_report(data, datetime(2026, 8, 30))
    workbook = load_workbook(BytesIO(report))

    details = workbook["Жилбыт"]
    rows = list(details.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][2] == "Аноним"
    assert rows[0][3] == "Аноним"
    assert "Секретное Имя" not in str(rows)


def test_report_shows_regular_user_data():
    ticket = Ticket(
        id=2,
        topic="Обычная тема",
        description="Текст",
        status=TicketStatus.NEW,
        is_anonymous=False,
    )
    ticket.user = User(full_name="Иван Иванов", dormitory="№2")
    ticket.department = Department(id=2, name="Информ")

    data = _make_data([ticket])
    report = build_daily_report(data, datetime(2026, 8, 30))
    workbook = load_workbook(BytesIO(report))

    details = workbook["Информ"]
    rows = list(details.iter_rows(min_row=2, values_only=True))
    assert rows[0][2] == "Иван Иванов"
    assert rows[0][3] == "№2"


def test_report_summary_counts_completed():
    tickets = []
    for i in range(3):
        t = Ticket(id=i + 1, topic=f"Тема {i}", status=TicketStatus.COMPLETED)
        t.department = Department(id=1, name="Жилбыт")
        tickets.append(t)

    data = _make_data(tickets)
    report = build_daily_report(data, datetime(2026, 8, 30))
    workbook = load_workbook(BytesIO(report))

    summary = workbook["Краткая информация"]
    rows = list(summary.iter_rows(min_row=2, values_only=True))
    total_row = next(r for r in rows if r[0] == "ИТОГО")
    assert total_row[8] == 3  # Всего
    assert total_row[5] == 3  # Выполнено
    assert total_row[9] == 100.0  # % выполнения


async def test_report_run_deduplication(db_session_maker):
    from core.reporting import is_report_already_sent, mark_report_sent

    day = date(2026, 8, 30)

    assert await is_report_already_sent(day) is False

    await mark_report_sent(day)
    assert await is_report_already_sent(day) is True

    # Повторная отметка не должна создавать дубликат
    await mark_report_sent(day)
    from sqlalchemy import select

    from core.models import ReportRun

    async with db_session_maker() as session:
        runs = (await session.scalars(select(ReportRun).where(ReportRun.report_date == day))).all()
        assert len(runs) == 1


async def test_send_report_to_vk_uploads_bytes_and_sends_document(monkeypatch):
    import os

    from core import reporting

    class FakeUploader:
        async def upload(self, file_source, peer_id, title):
            if isinstance(file_source, str) and os.path.exists(file_source):
                with open(file_source, "rb") as f:
                    assert f.read() == b"report-bytes"
            else:
                assert file_source == b"report-bytes"
            assert peer_id == 123
            assert title == "report.xlsx"
            return "doc-attachment"

    class FakeMessages:
        async def send(self, **kwargs):
            self.kwargs = kwargs

    class FakeApi:
        def __init__(self):
            self.messages = FakeMessages()

    monkeypatch.setattr(reporting, "DocMessagesUploader", lambda api: FakeUploader())
    api = FakeApi()

    await reporting.send_report_to_vk(api, 123, b"report-bytes", "report.xlsx")

    assert api.messages.kwargs["peer_id"] == 123
    assert api.messages.kwargs["attachment"] == "doc-attachment"
    assert api.messages.kwargs["message"] == "📊 Ежедневный отчёт."


async def test_start_report_scheduler_uses_current_loop(monkeypatch):
    from core import reporting

    async def fake_report_loop(api):
        return None

    monkeypatch.setattr(reporting, "_report_loop", fake_report_loop)
    task = reporting.start_report_scheduler(object())
    await task
    assert task.done()


def test_parse_report_date_valid():
    from core.reporting import parse_report_date

    assert parse_report_date("31.08.2026") == date(2026, 8, 31)
    assert parse_report_date("01.01.2025") == date(2025, 1, 1)
    assert parse_report_date("15.06.24") == date(2024, 6, 15)


def test_parse_report_date_invalid():
    from core.reporting import parse_report_date

    assert parse_report_date("2026-08-31") is None
    assert parse_report_date("31/08/2026") is None
    assert parse_report_date("не дата") is None
    assert parse_report_date("") is None
    assert parse_report_date(" 31.08.2026 ") == date(2026, 8, 31)  # whitespace stripped


async def test_get_report_for_period_requires_valid_range():
    from core.reporting import get_report_for_period

    date_from = date(2026, 9, 1)
    date_to = date(2026, 8, 31)  # раньше начала

    with pytest.raises(ValueError, match="не может быть позже"):
        await get_report_for_period(date_from, date_to)


async def test_get_report_for_period_collects_tickets(db_session_maker):
    from core.reporting import get_report_for_period

    async with db_session_maker() as session:
        user = User(vk_id=999)
        department = Department(name="Жилбыт")
        session.add_all([user, department])
        await session.flush()

        # Заявка в периоде
        ticket_in_period = Ticket(
            user_id=user.id,
            department_id=department.id,
            topic="Заявка в периоде",
            created_at=datetime(2026, 8, 15, 10, 0),
        )
        # Заявка до периода
        ticket_before = Ticket(
            user_id=user.id,
            department_id=department.id,
            topic="Заявка до периода",
            created_at=datetime(2026, 8, 1, 10, 0),
        )
        # Заявка после периода
        ticket_after = Ticket(
            user_id=user.id,
            department_id=department.id,
            topic="Заявка после периода",
            created_at=datetime(2026, 9, 1, 10, 0),
        )
        session.add_all([ticket_in_period, ticket_before, ticket_after])
        await session.commit()

    date_from = date(2026, 8, 10)
    date_to = date(2026, 8, 20)

    data = await get_report_for_period(date_from, date_to)
    assert len(data["tickets"]) == 1
    assert data["tickets"][0].topic == "Заявка в периоде"


async def test_get_report_for_period_leap_year(db_session_maker):
    """Отчёт за период, включающий 29 февраля високосного года, корректно собирает тикеты."""
    from core.reporting import get_report_for_period

    async with db_session_maker() as session:
        user = User(vk_id=888)
        department = Department(name="Жилбыт")
        session.add_all([user, department])
        await session.flush()

        ticket_leap = Ticket(
            user_id=user.id,
            department_id=department.id,
            topic="Тикет 29 февраля",
            created_at=datetime(2024, 2, 29, 12, 0),
        )
        session.add(ticket_leap)
        await session.commit()

    date_from = date(2024, 2, 28)
    date_to = date(2024, 3, 1)

    data = await get_report_for_period(date_from, date_to)
    assert len(data["tickets"]) == 1
    assert data["tickets"][0].topic == "Тикет 29 февраля"


async def test_get_report_for_period_full_year_window(db_session_maker):
    """12-месячное окно отчёта инклюзивно: включает крайние тикеты и отсекает внешние."""
    from core.reporting import get_report_for_period

    async with db_session_maker() as session:
        user = User(vk_id=777)
        department = Department(name="Информ")
        session.add_all([user, department])
        await session.flush()

        t_start = Ticket(
            user_id=user.id,
            department_id=department.id,
            topic="Тикет на старте года",
            created_at=datetime(2025, 1, 1, 0, 0, 5),
        )
        t_end = Ticket(
            user_id=user.id,
            department_id=department.id,
            topic="Тикет в конце года",
            created_at=datetime(2025, 12, 31, 23, 59, 50),
        )
        t_before = Ticket(
            user_id=user.id,
            department_id=department.id,
            topic="Тикет до начала года",
            created_at=datetime(2024, 12, 31, 23, 59, 59),
        )
        t_after = Ticket(
            user_id=user.id,
            department_id=department.id,
            topic="Тикет после окончания года",
            created_at=datetime(2026, 1, 1, 0, 0, 1),
        )
        session.add_all([t_start, t_end, t_before, t_after])
        await session.commit()

    data = await get_report_for_period(date(2025, 1, 1), date(2025, 12, 31))
    topics = {t.topic for t in data["tickets"]}
    assert "Тикет на старте года" in topics
    assert "Тикет в конце года" in topics
    assert "Тикет до начала года" not in topics
    assert "Тикет после окончания года" not in topics


async def test_get_report_for_period_empty_range_returns_clean_structure(db_session_maker):
    """При отсутствии тикетов в диапазоне возвращается пустой список и структура отделов."""
    from core.reporting import get_report_for_period

    data = await get_report_for_period(date(2020, 1, 1), date(2020, 1, 10))
    assert data["tickets"] == []
    assert isinstance(data["departments"], list)
